import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

import openpyxl
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.models import Income, Spending
import uuid


async def parse_excel_preview(content: bytes) -> dict:
    """Parse Excel file and return column headers + first 10 rows as preview."""
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"headers": [], "rows": [], "total_rows": 0}

    headers = [str(h) if h is not None else f"עמודה {i+1}" for i, h in enumerate(rows[0])]
    preview_rows = []
    for row in rows[1:11]:  # First 10 data rows
        preview_rows.append([str(cell) if cell is not None else "" for cell in row])

    wb.close()
    return {
        "headers": headers,
        "rows": preview_rows,
        "total_rows": len(rows) - 1,
    }


def _parse_date(value: Any) -> str:
    """Convert a date/datetime/string value to YYYY-MM-DD string."""
    if value is None:
        return date.today().isoformat()
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    try:
        return datetime.strptime(str(value), "%Y-%m-%d").date().isoformat()
    except ValueError:
        return date.today().isoformat()


async def parse_budget_excel(content: bytes) -> dict:
    """
    Parse a monthly budget Excel file with the fixed structure:
    - Sheet 1 (Summary): skipped
    - Sheets 2+: side-by-side expenses and income
        Expense cols (0-based): date=1, amount=2, notes=3, category_name=4
        Income cols (0-based):  date=6, amount=7, notes=8, category_name=9
    Data rows start at index 4 (rows 0-3 are headers/labels).
    Payment method: "cash" if notes contains "מזומן", else "card".
    """
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    sheet_names = wb.sheetnames

    # Skip the first sheet (Summary)
    data_sheets = sheet_names[1:]

    sheets = []
    all_category_names: set[str] = set()

    for sheet_name in data_sheets:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        expenses = []
        incomes = []

        for row in rows[4:]:  # Data starts at index 4
            # --- Expenses (cols B-E = indices 1-4) ---
            exp_date_raw = row[1] if len(row) > 1 else None
            exp_amount_raw = row[2] if len(row) > 2 else None
            exp_notes_raw = row[3] if len(row) > 3 else None
            exp_cat_raw = row[4] if len(row) > 4 else None

            if exp_amount_raw is not None:
                try:
                    exp_amount = float(Decimal(str(exp_amount_raw)))
                    if exp_amount > 0:
                        exp_notes = str(exp_notes_raw).strip() if exp_notes_raw else None
                        payment_method = "cash" if exp_notes and "מזומן" in exp_notes else "card"
                        cat_name = str(exp_cat_raw).strip() if exp_cat_raw else "שונות"
                        all_category_names.add(cat_name)
                        expenses.append({
                            "date": _parse_date(exp_date_raw),
                            "amount": exp_amount,
                            "notes": exp_notes,
                            "category_name": cat_name,
                            "payment_method": payment_method,
                        })
                except (InvalidOperation, ValueError):
                    pass

            # --- Income (cols G-J = indices 6-9) ---
            inc_date_raw = row[6] if len(row) > 6 else None
            inc_amount_raw = row[7] if len(row) > 7 else None
            inc_notes_raw = row[8] if len(row) > 8 else None
            inc_cat_raw = row[9] if len(row) > 9 else None

            if inc_amount_raw is not None:
                try:
                    inc_amount = float(Decimal(str(inc_amount_raw)))
                    if inc_amount > 0:
                        inc_notes = str(inc_notes_raw).strip() if inc_notes_raw else None
                        cat_name = str(inc_cat_raw).strip() if inc_cat_raw else "שונות"
                        all_category_names.add(cat_name)
                        incomes.append({
                            "date": _parse_date(inc_date_raw),
                            "amount": inc_amount,
                            "notes": inc_notes,
                            "category_name": cat_name,
                        })
                except (InvalidOperation, ValueError):
                    pass

        sheets.append({
            "name": sheet_name,
            "expenses": expenses,
            "incomes": incomes,
        })

    wb.close()
    return {
        "sheets": sheets,
        "all_category_names": sorted(all_category_names),
    }


async def import_spendings_bulk(
    db: AsyncSession,
    family_id: uuid.UUID,
    user_id: uuid.UUID,
    user_email: str,
    rows: list[dict[str, Any]],
    user_name: str | None = None,
) -> int:
    """
    Bulk insert spendings from imported rows.
    Each row must have: amount, category_id, date, payment_method, notes (optional)
    """
    count = 0
    for row in rows:
        try:
            amount_raw = row.get("amount")
            if amount_raw is None:
                continue
            amount = Decimal(str(amount_raw))

            category_id_raw = row.get("category_id")
            if not category_id_raw:
                continue
            category_id = uuid.UUID(str(category_id_raw))

            date_raw = row.get("date")
            if isinstance(date_raw, str):
                parsed_date = datetime.strptime(date_raw, "%Y-%m-%d").date()
            elif isinstance(date_raw, date):
                parsed_date = date_raw
            else:
                parsed_date = date.today()

            payment_method = row.get("payment_method", "cash")
            card_id_raw = row.get("card_id")
            card_id = uuid.UUID(str(card_id_raw)) if card_id_raw else None
            notes = row.get("notes")

            spending = Spending(
                family_id=family_id,
                user_id=user_id,
                user_email=user_email,
                user_name=user_name,
                amount=amount,
                category_id=category_id,
                payment_method=payment_method,
                card_id=card_id,
                date=parsed_date,
                notes=notes,
            )
            db.add(spending)
            count += 1
        except (ValueError, InvalidOperation, AttributeError):
            continue

    if count > 0:
        await db.commit()
    return count


async def import_incomes_bulk(
    db: AsyncSession,
    family_id: uuid.UUID,
    user_id: uuid.UUID,
    user_email: str,
    rows: list[dict[str, Any]],
    user_name: str | None = None,
) -> int:
    """
    Bulk insert incomes from imported rows.
    Each row must have: amount, date; category_id and notes are optional.
    """
    count = 0
    for row in rows:
        try:
            amount_raw = row.get("amount")
            if amount_raw is None:
                continue
            amount = Decimal(str(amount_raw))

            date_raw = row.get("date")
            if isinstance(date_raw, str):
                parsed_date = datetime.strptime(date_raw, "%Y-%m-%d").date()
            elif isinstance(date_raw, date):
                parsed_date = date_raw
            else:
                parsed_date = date.today()

            category_id_raw = row.get("category_id")
            category_id = uuid.UUID(str(category_id_raw)) if category_id_raw else None

            notes = row.get("notes")

            income = Income(
                family_id=family_id,
                user_id=user_id,
                user_email=user_email,
                user_name=user_name,
                amount=amount,
                category_id=category_id,
                date=parsed_date,
                notes=notes,
            )
            db.add(income)
            count += 1
        except (ValueError, InvalidOperation, AttributeError):
            continue

    if count > 0:
        await db.commit()
    return count
